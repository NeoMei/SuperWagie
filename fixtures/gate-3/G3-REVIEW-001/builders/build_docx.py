from __future__ import annotations

import copy
import argparse
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_run_font(run, name='Aptos', size=11, bold=None, color='243447'):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for edge, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths_dxa, total_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_w.set(qn('w:w'), str(total_dxa))
    tbl_ind = tbl_pr.first_child_found_in('w:tblInd')
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:type'), 'dxa')
    tbl_ind.set(qn('w:w'), '120')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in('w:tcW')
            tc_w.set(qn('w:type'), 'dxa')
            tc_w.set(qn('w:w'), str(widths_dxa[idx]))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(paragraph.add_run('Page '), size=9, color='667085')
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t')
    txt.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run = paragraph.add_run()._r
    run.extend([begin, instr, separate, txt, end])
    set_run_font(paragraph.add_run(' of 30'), size=9, color='667085')


def configure_section(section, label, landscape=False):
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.28)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Inches(11), Inches(8.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.clear()
    set_run_font(hp.add_run(f'SuperWagie Office Reviewer PoC  |  {label}'), size=9, bold=True, color='375A7F')
    fp = section.footer.paragraphs[0]
    fp.clear()
    add_page_field(fp)


def configure_styles(doc):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ('Title', 25, '17324D', 0, 10),
        ('Subtitle', 13, '52677B', 0, 18),
        ('Heading 1', 16, '2E74B5', 14, 7),
        ('Heading 2', 13, '2E74B5', 10, 5),
        ('Heading 3', 12, '1F4D78', 8, 4),
    ]:
        style = styles[name]
        style.font.name = 'Aptos'
        style._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Aptos')
        style._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Aptos')
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if 'TOC Fixture' not in [s.name for s in styles]:
        toc = styles.add_style('TOC Fixture', WD_STYLE_TYPE.PARAGRAPH)
        toc.base_style = styles['Normal']
        toc.font.name = 'Aptos'
        toc.font.size = Pt(10)
        toc.paragraph_format.space_after = Pt(2)


def link_heading_numbering(doc):
    numbering = doc.part.numbering_part.element
    abs_ids = [int(v) for v in numbering.xpath('./w:abstractNum/@w:abstractNumId')]
    num_ids = [int(v) for v in numbering.xpath('./w:num/@w:numId')]
    abstract_id = max(abs_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    multi = OxmlElement('w:multiLevelType')
    multi.set(qn('w:val'), 'multilevel')
    abstract.append(multi)
    for level, (fmt, text, style_name, left, hanging) in enumerate([
        ('decimal', '%1', 'Heading1', 0, 0),
        ('decimal', '%1.%2', 'Heading2', 360, 180),
        ('decimal', '%1.%2.%3', 'Heading3', 720, 180),
    ]):
        lvl = OxmlElement('w:lvl')
        lvl.set(qn('w:ilvl'), str(level))
        start = OxmlElement('w:start'); start.set(qn('w:val'), '1'); lvl.append(start)
        num_fmt = OxmlElement('w:numFmt'); num_fmt.set(qn('w:val'), fmt); lvl.append(num_fmt)
        p_style = OxmlElement('w:pStyle'); p_style.set(qn('w:val'), style_name); lvl.append(p_style)
        lvl_text = OxmlElement('w:lvlText'); lvl_text.set(qn('w:val'), text); lvl.append(lvl_text)
        suff = OxmlElement('w:suff'); suff.set(qn('w:val'), 'space'); lvl.append(suff)
        p_pr = OxmlElement('w:pPr')
        tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab'); tab.set(qn('w:val'), 'num'); tab.set(qn('w:pos'), str(left + 360)); tabs.append(tab); p_pr.append(tabs)
        ind = OxmlElement('w:ind'); ind.set(qn('w:left'), str(left + 360)); ind.set(qn('w:hanging'), str(hanging or 360)); p_pr.append(ind)
        lvl.append(p_pr)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(num_id))
    abs_ref = OxmlElement('w:abstractNumId'); abs_ref.set(qn('w:val'), str(abstract_id)); num.append(abs_ref)
    numbering.append(num)
    for style_name, level in [('Heading 1', 0), ('Heading 2', 1), ('Heading 3', 2)]:
        p_pr = doc.styles[style_name]._element.get_or_add_pPr()
        old = p_pr.find(qn('w:numPr'))
        if old is not None:
            p_pr.remove(old)
        num_pr = OxmlElement('w:numPr')
        ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), str(level)); num_pr.append(ilvl)
        num_id_el = OxmlElement('w:numId'); num_id_el.set(qn('w:val'), str(num_id)); num_pr.append(num_id_el)
        p_pr.append(num_pr)
    return num_id


def add_toc_field(paragraph):
    run = paragraph.add_run()._r
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin'); begin.set(qn('w:dirty'), 'true')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    cached = OxmlElement('w:t'); cached.text = 'Field-driven TOC: update fields in WPS/Word'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run.extend([begin, instr, separate, cached, end])


def add_table(doc, landscape=False):
    values = [
        ['Case', 'Expected reviewer behavior', 'Risk', 'Status'],
        ['Pagination', 'Keep page identity stable after content reflow', 'High', 'Ready'],
        ['Typography', 'Expose substitution without losing text', 'Medium', 'Ready'],
        ['Annotations', 'Relocate by revision, page, bbox, and content hash', 'High', 'Ready'],
    ]
    table = doc.add_table(rows=len(values), cols=4)
    table.style = 'Table Grid'
    widths = [1500, 4200 if not landscape else 7600, 1200, 2460]
    if landscape:
        widths = [1800, 6800, 1700, 2660]
        total = 12960
    else:
        total = 9360
    set_table_geometry(table, widths, total)
    for r, row in enumerate(table.rows):
        if r == 0:
            set_repeat_table_header(row)
        for c, cell in enumerate(row.cells):
            cell.text = values[r][c]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=9.5, bold=(r == 0), color='FFFFFF' if r == 0 else '243447')
            if r == 0:
                shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), '375A7F'); cell._tc.get_or_add_tcPr().append(shd)
    return table


def add_floating_picture(paragraph, path):
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(1.55), height=Inches(0.78))
    drawing = run._r.find(qn('w:drawing'))
    inline = drawing.find(qn('wp:inline'))
    anchor = OxmlElement('wp:anchor')
    for key, value in {
        'distT': '0', 'distB': '0', 'distL': '114300', 'distR': '114300',
        'simplePos': '0', 'relativeHeight': '251658240', 'behindDoc': '0',
        'locked': '0', 'layoutInCell': '1', 'allowOverlap': '0'
    }.items():
        anchor.set(key, value)
    simple = OxmlElement('wp:simplePos'); simple.set('x', '0'); simple.set('y', '0'); anchor.append(simple)
    pos_h = OxmlElement('wp:positionH'); pos_h.set('relativeFrom', 'column')
    pos_h_off = OxmlElement('wp:posOffset'); pos_h_off.text = '4200000'; pos_h.append(pos_h_off); anchor.append(pos_h)
    pos_v = OxmlElement('wp:positionV'); pos_v.set('relativeFrom', 'paragraph')
    pos_v_off = OxmlElement('wp:posOffset'); pos_v_off.text = '120000'; pos_v.append(pos_v_off); anchor.append(pos_v)
    for tag in ['wp:extent', 'wp:effectExtent']:
        node = inline.find(qn(tag))
        if node is not None:
            anchor.append(copy.deepcopy(node))
    wrap = OxmlElement('wp:wrapSquare'); wrap.set('wrapText', 'bothSides'); anchor.append(wrap)
    for tag in ['wp:docPr', 'wp:cNvGraphicFramePr', 'a:graphic']:
        node = inline.find(qn(tag))
        if node is not None:
            anchor.append(copy.deepcopy(node))
    drawing.replace(inline, anchor)


def add_body(doc, text, *, missing_font=False, cjk=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    name = 'SuperWagieMissingFont-927' if missing_font else ('Arial Unicode MS' if cjk else 'Aptos')
    set_run_font(r, name=name, size=11)
    return p


def make_asset(asset_path: Path):
    image = Image.new('RGB', (900, 420), '#EAF2F8')
    d = ImageDraw.Draw(image)
    d.rounded_rectangle((22, 22, 878, 398), radius=28, outline='#2E74B5', width=9, fill='#F7FBFF')
    for i, color in enumerate(['#2E74B5', '#4DA3D9', '#7CC4E4', '#A9D8EA']):
        d.rectangle((70 + i * 180, 90, 190 + i * 180, 320 - i * 28), fill=color)
    d.text((70, 345), 'FLOATING IMAGE / WRAP TEST', fill='#17324D')
    image.save(asset_path)


def main(output: Path, work_dir: Path):
    output = output.resolve()
    work_dir = work_dir.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    work_dir.mkdir(parents=True, exist_ok=True)
    asset = work_dir / 'docx-floating-test.png'
    make_asset(asset)
    doc = Document()
    configure_styles(doc)
    link_heading_numbering(doc)
    configure_section(doc.sections[0], 'Portrait A', landscape=False)
    doc.core_properties.title = 'SuperWagie Reviewer Torture Fixture - 30 Pages'
    doc.core_properties.subject = 'Project-owned deterministic Office review fixture'
    doc.core_properties.author = 'SuperWagie Project Fixture Builder'

    # Page 1 - cover + true TOC field + deterministic cached/static entries.
    p = doc.add_paragraph(style='Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run('Office Reviewer Torture Fixture'), size=25, bold=True, color='17324D')
    p = doc.add_paragraph(style='Subtitle'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run('30-page deterministic DOCX for pagination, structure, and fidelity validation'), size=13, color='52677B')
    h = doc.add_paragraph('Table of Contents', style='Heading 1')
    toc_field = doc.add_paragraph(style='TOC Fixture'); add_toc_field(toc_field)
    for entry in [
        '1. Fixture overview ........................................ 2',
        '2. Mixed content and notes ................................. 5',
        '3. Landscape stress section ............................... 10',
        '4. Anchored images and repagination ....................... 14',
        '5. Typography substitution ................................ 24',
        '6. Final acceptance page .................................. 30',
    ]:
        doc.add_paragraph(entry, style='TOC Fixture')
    doc.add_page_break()

    titles = {
        2: 'Fixture overview', 3: 'Linked numbering hierarchy', 4: 'Chinese and Latin shaping',
        5: 'True footnote boundary', 6: 'Dense paragraph rhythm', 7: 'Native table geometry',
        8: 'Header and footer continuity', 9: 'Portrait section boundary', 10: 'Landscape section begins',
        11: 'Wide table in landscape', 12: 'Landscape section ends', 13: 'Portrait resumes',
        14: 'Floating image and wrap', 15: 'Image-following content', 16: 'Pre-edit pagination baseline',
        17: 'Repagination edit stress', 18: 'Post-edit flow', 19: 'Nested heading continuation',
        20: 'Review annotation anchors', 21: 'Revision identity', 22: 'Long Latin tokens',
        23: 'Chinese punctuation and spacing', 24: 'Missing font substitution', 25: 'Mixed-script fallback',
        26: 'Late-document table', 27: 'Footer field continuity', 28: 'Final reflow boundary',
        29: 'Acceptance preflight', 30: 'Fixture completion',
    }

    for page in range(2, 31):
        if page == 10:
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
            configure_section(sec, 'Landscape', landscape=True)
        elif page == 13:
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
            configure_section(sec, 'Portrait B', landscape=False)

        doc.add_paragraph(titles[page], style='Heading 1')
        if page in (3, 19):
            doc.add_paragraph('A linked secondary level', style='Heading 2')
            doc.add_paragraph('A linked tertiary level', style='Heading 3')
        add_body(doc, f'Page {page} is project-owned fixture content. It is deliberately concise so the explicit page boundary remains deterministic across bundled LibreOffice rendering and later real-WPS truth rendering.')
        if page == 4:
            add_body(doc, '中文排版测试：超级牛马审阅视图需要保持分页、编号、表格、字体与批注定位。 Latin shaping test: Office Reviewer preserves revision identity and editable structure.', cjk=True)
        elif page == 5:
            add_body(doc, 'This sentence owns a true footnote reference [[FN]] and remains editable after the note part is inserted.')
        elif page == 7:
            add_table(doc)
        elif page == 8:
            add_body(doc, 'The running header changes by section; the footer contains a real PAGE field and a fixed total of 30.')
        elif page == 11:
            add_table(doc, landscape=True)
        elif page == 14:
            anchor_p = doc.add_paragraph()
            add_floating_picture(anchor_p, asset)
            add_body(doc, 'Text continues beside the anchored image. The picture is a true wp:anchor with square wrapping, not an inline-only drawing.')
        elif page == 16:
            add_body(doc, 'Short baseline.')
        elif page == 17:
            add_body(doc, 'REPAGINATION_BASELINE')
            for idx in range(1, 5):
                add_body(doc, f'Flow paragraph {idx}: the accepted insertion is long enough to change line wrapping and force downstream pagination calculations while the fixture retains its explicit terminal page count.')
        elif page == 20:
            add_body(doc, 'Anchor token: review-semantic-object-20 / content-hash-seed-20 / bbox-fallback-required.')
        elif page == 22:
            add_body(doc, 'LongToken_' + 'ReviewerReflowBoundary' * 5 + ' tests wrapping without corruption.')
        elif page == 23:
            add_body(doc, '中文标点测试：“批注、差异、接受／拒绝、版本确认”；全角括号（稳定）与英文 parentheses (stable) 并存。', cjk=True)
        elif page == 24:
            add_body(doc, 'MISSING FONT SUBSTITUTION TEST - readable fallback must preserve this text.', missing_font=True)
        elif page == 25:
            add_body(doc, '混合脚本 Mixed-script fallback: ABC 123 审阅 页面 Revision Ω.', missing_font=True, cjk=True)
        elif page == 26:
            add_table(doc)
        elif page == 28:
            add_body(doc, 'A late edit here should invalidate cached previews whose cache key includes artifact hash, renderer version, font environment, and parameters.')
        elif page == 30:
            add_body(doc, 'END OF 30-PAGE FIXTURE. Expected rendered page count: 30. Primary verification is structural plus every-page PNG inspection; later Gate tasks use real WPS as the visual fact source.')

        if page < 30 and page not in (9, 12):
            doc.add_page_break()

    # Deterministic metadata and explicit field-update hint.
    settings = doc.settings.element
    update_fields = settings.find(qn('w:updateFields'))
    if update_fields is None:
        update_fields = OxmlElement('w:updateFields')
        settings.append(update_fields)
    update_fields.set(qn('w:val'), 'true')
    doc.save(output)
    print(output)


def parse_args():
    parser = argparse.ArgumentParser(description='Build the project-owned DOCX torture fixture base package.')
    parser.add_argument('--output', type=Path, default=Path.cwd() / 'reviewer-torture-base.docx')
    parser.add_argument('--work-dir', type=Path, default=Path.cwd() / '.build-docx')
    return parser.parse_args()


if __name__ == '__main__':
    args = parse_args()
    main(args.output, args.work_dir)
